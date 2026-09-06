import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY_DARK = RGBColor(15, 23, 42)      # #0f172a
    C_NAVY_CARD = RGBColor(30, 41, 59)      # #1e293b
    C_BLUE_PRIMARY = RGBColor(2, 132, 199)  # #0284c7
    C_TEAL_ACCENT = RGBColor(13, 148, 136)  # #0d9488
    C_CYAN_TEXT = RGBColor(56, 189, 248)    # #38bdf8
    C_WHITE = RGBColor(255, 255, 255)
    C_MUTED = RGBColor(148, 163, 184)       # #94a3b8
    C_CARD_BORDER = RGBColor(51, 65, 85)    # #334155
    C_GREEN = RGBColor(34, 197, 94)         # #22c55e
    C_RED_ACCENT = RGBColor(239, 68, 68)    # #ef4444
    C_AMBER = RGBColor(245, 158, 11)        # #f59e0b

    def add_header(slide, title_text, category_text="SMART INDIA HACKATHON 2026 | SOFTWARE EDITION"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_sub = tf.paragraphs[0]
        p_sub.text = category_text.upper()
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = C_CYAN_TEXT
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_WHITE

    def add_card(slide, left, top, width, height, bg_color=C_NAVY_CARD, border_color=C_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY_DARK
    bg1.line.fill.background()

    card1 = add_card(s1, Inches(1.0), Inches(0.85), Inches(11.333), Inches(5.8), C_NAVY_CARD, C_BLUE_PRIMARY)
    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.6)
    tf1.margin_right = Inches(0.6)
    tf1.margin_top = Inches(0.45)

    p = tf1.paragraphs[0]
    p.text = "SMART INDIA HACKATHON (SIH) 2026 — OFFICIAL IDEA PRESENTATION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_CYAN_TEXT
    
    p = tf1.add_paragraph()
    p.text = "MahaHealth Connect"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(6)

    p = tf1.add_paragraph()
    p.text = "Dual-Outlet Rural Health Infrastructure & Voice-First Multilingual Tele-Platform"
    p.font.size = Pt(17)
    p.font.color.rgb = C_CYAN_TEXT
    p.space_after = Pt(20)

    info_box = s1.shapes.add_textbox(Inches(1.6), Inches(3.4), Inches(10.133), Inches(2.8))
    tf_info = info_box.text_frame
    tf_info.word_wrap = True

    p = tf_info.paragraphs[0]
    p.text = "📌 Problem Statement ID: 26133"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(4)

    p = tf_info.add_paragraph()
    p.text = "🏛️ Organization: Government of Maharashtra (Maharashtra State Innovation Society)"
    p.font.size = Pt(13)
    p.font.color.rgb = C_MUTED
    p.space_after = Pt(4)

    p = tf_info.add_paragraph()
    p.text = "🏷️ Theme: MedTech / BioTech / HealthTech  |  Category: Software"
    p.font.size = Pt(13)
    p.font.color.rgb = C_MUTED
    p.space_after = Pt(14)

    p = tf_info.add_paragraph()
    p.text = "🚀 Innovation: Village Dual-Outlet Model • Illiterate-Friendly Voice/Pictorial WebApp • 1-on-1 Agent Guidance • Direct SOS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    # =========================================================================
    # SLIDE 2: GROUND REALITIES & PROBLEM ANALYSIS
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    bg2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = C_NAVY_DARK
    bg2.line.fill.background()
    add_header(s2, "Problem Statement & Rural Healthcare Bottlenecks")

    cards_data_s2 = [
        ("1. The Rural Access & Literacy Barrier", 
         "Over 45% of villagers in tribal belts (Gadchiroli, Nandurbar, Melghat) are illiterate or non-tech savvy. Conventional healthcare apps with dense English/complex forms completely fail to serve them.",
         C_RED_ACCENT),
        ("2. The Dual Friction: Care vs Supplies", 
         "Inner villages lack doctors for early triage. Meanwhile, diagnostic labs and pharmacies are located far away in cities, forcing villagers to travel 40-80 km just to collect medicines or test reports.",
         C_CYAN_TEXT),
        ("3. Panic & Emergency Isolation", 
         "During acute emergencies (cardiac arrest, snakebites, maternal hemorrhages), villagers have no direct panic button or 1-on-1 human guidance to direct ambulances or perform life-saving first aid.",
         C_RED_ACCENT),
        ("4. Broken Doctor-Patient Continuity", 
         "Patients travel hours only to find specialist doctors unavailable. Past test reports on paper get destroyed or misplaced, preventing longitudinal care across sub-centres and district hospitals.",
         C_CYAN_TEXT)
    ]

    coords = [
        (Inches(0.8), Inches(1.6)),
        (Inches(6.9), Inches(1.6)),
        (Inches(0.8), Inches(4.45)),
        (Inches(6.9), Inches(4.45))
    ]

    for idx, (title, desc, border_col) in enumerate(cards_data_s2):
        l, t = coords[idx]
        c = add_card(s2, l, t, Inches(5.6), Inches(2.55), C_NAVY_CARD, border_col)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 3: PROPOSED ARCHITECTURE — DUAL-OUTLET MODEL + WEBAPP
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    bg3 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = C_NAVY_DARK
    bg3.line.fill.background()
    add_header(s3, "Proposed Solution: Dual Physical Outlets + Multilingual WebApp")

    pillars = [
        ("📍 Outlet 1: Gram Care Point\n(Inside Village Centre)", 
         "• Located at village core (Gram Panchayat/Chhavadi)\n• Dedicated touchpoint for patient queries & digital intake\n• Assisted check-ins by trained local health volunteers\n• Routine vitals monitoring (BP, SpO2, Temperature)",
         C_CYAN_TEXT),
        ("📍 Outlet 2: Common Diagnostic Hub\n(Outskirts Junction)", 
         "• Situated on arterial road serving a 3-5 village cluster\n• Houses cold-chain medical supplies & bulk inventory\n• Sample collection & instant diagnostic report hub\n• Staging base for 108 emergency ambulances & oxygen",
         C_TEAL_ACCENT),
        ("🌐 Multilingual Universal WebApp\n(Illiterate-Friendly UX)", 
         "• 100% Pictorial & Color-Coded Icon Navigation\n• Multilingual Voice Assistant (Marathi, Gondi, Hindi)\n• High-contrast, zero-typing audio-guided workflows\n• Works seamlessly on any basic smartphone or tablet",
         C_BLUE_PRIMARY),
        ("🤝 1-on-1 Guided Care & Direct SOS\n(Human Assisted Triage)", 
         "• One-Touch Emergency SOS: Direct 108 ambulance dispatch\n• 1-on-1 Live Agent Video/Audio guidance for non-literate users\n• Direct Doctor Tele-Consultations with auto-routed prescriptions to the Outskirts Supply Hub",
         C_GREEN)
    ]

    for idx, (title, desc, accent) in enumerate(pillars):
        l = Inches(0.8 + idx * 2.95)
        t = Inches(1.5)
        c = add_card(s3, l, t, Inches(2.7), Inches(5.4), C_NAVY_CARD, accent)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.22)
        tf.margin_right = Inches(0.22)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_CYAN_TEXT
        p.space_after = Pt(12)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 4: TECHNICAL ARCHITECTURE & FLOW
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    bg4 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg4.fill.solid()
    bg4.fill.fore_color.rgb = C_NAVY_DARK
    bg4.line.fill.background()
    add_header(s4, "End-to-End System Workflow & Technical Architecture")

    layers = [
        ("Layer 1: Frontline Village Outlets & Assistive WebApp", 
         "• Tech: Responsive React/Flutter PWA with WebAudio API & Canvas Icon System\n• Accessibility: Voice-guided navigation via Bhashini Speech-to-Text & Text-to-Speech in Marathi/Hindi\n• Offline Cache: Local IndexedDB / SQLite caching for low-network resilience",
         Inches(1.5)),
        ("Layer 2: Real-Time Communication & Telehealth Engine", 
         "• 1-on-1 Agent & Doctor Video/Audio: LiveKit / WebRTC optimized with Opus codec down to 15 kbps\n• Direct SOS Gateway: WebSocket emergency dispatch with GPS broadcast to 108 and Outskirts Hubs\n• Backend Services: High-throughput FastAPI (Python) & Node.js microservices with Redis queues",
         Inches(3.35)),
        ("Layer 3: Diagnostic, Inventory Ledger & ABDM Standards", 
         "• Inventory Sync: Real-time supply tracking between Outskirts Diagnostic Hub and District Warehouse\n• Health Records: HL7 / FHIR R4 interoperable JSON schemas linked to patient ABHA IDs\n• Security: DPDP Act 2023 compliant data encryption (AES-256 at rest, TLS 1.3 in transit)",
         Inches(5.2))
    ]

    for title, desc, top_pos in layers:
        c = add_card(s4, Inches(0.8), top_pos, Inches(11.733), Inches(1.65), C_NAVY_CARD, C_CARD_BORDER)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.2)
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_CYAN_TEXT
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 5: IMPLEMENTATION & FEASIBILITY ROADMAP
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    bg5 = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg5.fill.solid()
    bg5.fill.fore_color.rgb = C_NAVY_DARK
    bg5.line.fill.background()
    add_header(s5, "Phased Implementation & Village Cluster Rollout Plan")

    phases = [
        ("Phase 1: Pilot Cluster (M1–M3)",
         "📍 Target: 10 Villages in Gadchiroli / Melghat\n• Establish 10 Inner Village Care Points & 2 Outskirts Diagnostic Hubs\n• Deploy multilingual voice WebApp with local Marathi/Gondi audio prompts\n• Train 30 village health agents for 1-on-1 user onboarding.",
         C_BLUE_PRIMARY),
        ("Phase 2: District Scaling (M4–M8)",
         "📍 Target: 5 High-Priority Rural Districts\n• Deploy 100 Outskirts Logistics Hubs serving 400+ villages\n• Full integration with 108 Emergency Medical Services & District Civil Hospitals\n• Activate automated diagnostic report push to patient WhatsApp/SMS.",
         C_TEAL_ACCENT),
        ("Phase 3: Statewide Network (M9–M12+)",
         "📍 Target: All 36 Districts of Maharashtra\n• Full federation with Ayushman Bharat Digital Mission (ABDM)\n• AI-driven supply reorder prediction based on seasonal outbreak patterns\n• Standard operating framework for pan-India state adoption.",
         C_GREEN)
    ]

    for idx, (title, desc, col) in enumerate(phases):
        l = Inches(0.8 + idx * 3.95)
        t = Inches(1.5)
        c = add_card(s5, l, t, Inches(3.7), Inches(5.3), C_NAVY_CARD, col)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_after = Pt(14)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 6: MEASURABLE IMPACT & UNFAIR ADVANTAGES
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    bg6 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg6.fill.solid()
    bg6.fill.fore_color.rgb = C_NAVY_DARK
    bg6.line.fill.background()
    add_header(s6, "Measurable Social Impact & Key Differentiators")

    c_left = add_card(s6, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3), C_NAVY_CARD, C_GREEN)
    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_right = Inches(0.3)
    tf_l.margin_top = Inches(0.3)

    p = tf_l.paragraphs[0]
    p.text = "📊 Key Measurable Impact"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GREEN
    p.space_after = Pt(14)

    metrics = [
        ("70% Reduction in Diagnostic Travel", "Villagers obtain medicines and tests at the Outskirts Hub without journeying 60km to cities."),
        ("100% Inclusive for Illiterate Citizens", "Voice-guided icon navigation removes all reading/writing barriers for elderly and tribal villagers."),
        ("< 15 Minute Emergency Response", "Direct SOS routing directly alerts nearby ambulances and staging hubs."),
        ("Zero Lost Prescriptions", "Digital prescription auto-sent to Outskirts Hub for immediate dispensing.")
    ]
    for m_title, m_desc in metrics:
        p = tf_l.add_paragraph()
        p.text = f"• {m_title}: {m_desc}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_WHITE
        p.space_after = Pt(8)

    c_right = add_card(s6, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3), C_NAVY_CARD, C_CYAN_TEXT)
    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_right = Inches(0.3)
    tf_r.margin_top = Inches(0.3)

    p = tf_r.paragraphs[0]
    p.text = "🏆 Why Our Solution Wins"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CYAN_TEXT
    p.space_after = Pt(14)

    diffs = [
        ("Dual-Outlet Micro-Infrastructure", "Splits administrative query handling inside the village from high-throughput supply logistics on the outskirts."),
        ("Illiterate-First Voice UX", "Replaces complex text forms with spoken Marathi/Hindi/Gondi prompts and clear pictorial visual cues."),
        ("1-on-1 Human Agent Fallback", "Provides instant live human support when an elderly or panicked villager needs guidance."),
        ("End-to-End Closed Loop", "Direct bridge between Doctor Tele-Consult ➔ Outskirts Medicine Dispensing ➔ Follow-up.")
    ]
    for d_title, d_desc in diffs:
        p = tf_r.add_paragraph()
        p.text = f"⭐ {d_title}\n  {d_desc}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_WHITE
        p.space_after = Pt(8)

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out = r"C:\Users\Lenovo\Downloads\SIH2026_PS_26133_Idea_Presentation.pptx"
    create_deck(out)
    scratch_out = r"C:\Users\Lenovo\.gemini\antigravity\scratch\mahahealth_connect\SIH2026_PS_26133_Idea_Presentation.pptx"
    create_deck(scratch_out)
